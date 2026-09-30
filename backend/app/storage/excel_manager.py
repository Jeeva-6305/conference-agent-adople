import logging
import os
import re
from datetime import date
from typing import List, Optional
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..config import settings
from ..models import Conference, Speaker

logger = logging.getLogger(__name__)

EXCEL_COLUMNS = [
    "Conference ID",
    "Conference Title",
    "Industry / Category",
    "Conference Start Date",
    "Conference End Date",
    "Speakers Available",
    "Venue",
    "City",
    "Country",
    "Organizer",
    "Official Conference URL",
    "Registration URL",
    "Source URL",
    "Publication Date"
]

SPEAKER_COLUMNS = [
    "Speaker Name",
    "Conference Name",
    "Company / Organization",
    "Job Role / Designation",
    "Company Name",
    "Official Company Website URL",
    "Speaker LinkedIn Profile URL",
    "Location"
]

class ExcelManager:
    def __init__(self, file_path: str = settings.EXCEL_OUTPUT_PATH):
        self.file_path = file_path
        self._ensure_file_exists()

    def _ensure_file_exists(self):
        if not os.path.exists(self.file_path):
            self.reset_excel()

    def reset_excel(self):
        """Creates or resets the Excel file with two styled sheets: USA Conferences & Speakers."""
        wb = Workbook()
        
        # 1. Main Sheet: USA Conferences
        ws_conf = wb.active
        ws_conf.title = "USA Conferences"
        
        conf_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        conf_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws_conf.append(EXCEL_COLUMNS)
        ws_conf.row_dimensions[1].height = 28

        for col_idx, col_name in enumerate(EXCEL_COLUMNS, 1):
            cell = ws_conf.cell(row=1, column=col_idx)
            cell.fill = conf_header_fill
            cell.font = conf_header_font
            cell.alignment = align_center

        self._auto_fit_columns(ws_conf)

        # 2. Second Sheet: Speakers
        ws_spk = wb.create_sheet(title="Speakers")
        spk_header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        spk_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        ws_spk.append(SPEAKER_COLUMNS)
        ws_spk.row_dimensions[1].height = 28

        for col_idx, col_name in enumerate(SPEAKER_COLUMNS, 1):
            cell = ws_spk.cell(row=1, column=col_idx)
            cell.fill = spk_header_fill
            cell.font = spk_header_font
            cell.alignment = align_center

        self._auto_fit_columns(ws_spk)

        self._safe_save(wb)
        logger.info(f"Initialized blank Excel workbook with 'USA Conferences' and 'Speakers' sheets at {self.file_path}")

    def _safe_save(self, wb) -> bool:
        try:
            wb.save(self.file_path)
            return True
        except PermissionError:
            logger.warning(f"File {self.file_path} is locked by an external application (like Excel). Retrying...")
            import time
            for _ in range(3):
                time.sleep(0.5)
                try:
                    wb.save(self.file_path)
                    return True
                except PermissionError:
                    pass
            logger.warning(f"Could not overwrite {self.file_path} because it is currently locked in Excel.")
            return False

    def _auto_fit_columns(self, ws):
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
            for cell in col:
                if cell.row > 1:
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")

    def sync_published_conferences(self, conferences: List[Conference], speakers: Optional[List[Speaker]] = None) -> int:
        """
        Completely refreshes the Excel workbook with genuine verified conferences
        and their associated speakers sheet.
        """
        self.reset_excel()
        return self.publish_conferences(conferences, speakers)

    def publish_conferences(self, conferences: List[Conference], speakers: Optional[List[Speaker]] = None) -> int:
        """
        Appends newly due conferences into the 'USA Conferences' sheet
        and their detailed speaker profiles into the 'Speakers' sheet.
        """
        if not conferences:
            return 0

        self._ensure_file_exists()
        wb = load_workbook(self.file_path)
        
        if "USA Conferences" not in wb.sheetnames:
            ws_conf = wb.active
            ws_conf.title = "USA Conferences"
        else:
            ws_conf = wb["USA Conferences"]

        if "Speakers" not in wb.sheetnames:
            ws_spk = wb.create_sheet(title="Speakers")
            ws_spk.append(SPEAKER_COLUMNS)
        else:
            ws_spk = wb["Speakers"]

        # Collect existing conference IDs and titles
        existing_ids = set()
        existing_titles = set()
        for row in ws_conf.iter_rows(min_row=2, max_col=2, values_only=True):
            if row[0]:
                existing_ids.add(str(row[0]).strip())
            if len(row) > 1 and row[1]:
                existing_titles.add(str(row[1]).strip().lower())

        # Collect existing speakers (speaker_name, conf_title)
        existing_speakers = set()
        for row in ws_spk.iter_rows(min_row=2, max_col=2, values_only=True):
            if row[0] and len(row) > 1 and row[1]:
                existing_speakers.add((str(row[0]).strip().lower(), str(row[1]).strip().lower()))

        added_conf_count = 0
        added_spk_count = 0

        for conf in conferences:
            title_clean = conf.conference_title.strip().lower()
            if conf.conference_id in existing_ids or title_clean in existing_titles:
                continue

            # Determine whether speakers/auditors are available (Yes/No)
            speakers_available_str = conf.speakers_available or "Yes"
            
            # Publication Date: MUST be the original conference publication date
            pub_date_str = conf.publication_date.isoformat() if conf.publication_date else ""

            row_data = [
                conf.conference_id,
                conf.conference_title,
                conf.industry_category,
                conf.start_date.isoformat() if conf.start_date else "",
                conf.end_date.isoformat() if conf.end_date else "",
                speakers_available_str,
                conf.venue,
                conf.city,
                conf.country,
                conf.organizer,
                conf.official_conference_url,
                conf.registration_url,
                conf.source_url,
                pub_date_str
            ]
            ws_conf.append(row_data)
            added_conf_count += 1
            existing_ids.add(conf.conference_id)
            existing_titles.add(title_clean)

        # Write explicit Speaker models into the Speakers sheet
        if speakers:
            for spk in speakers:
                key = (spk.speaker_name.strip().lower(), spk.conference_name.strip().lower())
                if key not in existing_speakers:
                    ws_spk.append([
                        spk.speaker_name,
                        spk.conference_name,
                        spk.company_organization or spk.company_name,
                        spk.job_role_designation,
                        spk.company_name,
                        getattr(spk, "official_company_website_url", "") or "",
                        getattr(spk, "linkedin_url", "") or "",
                        spk.location
                    ])
                    existing_speakers.add(key)
                    added_spk_count += 1

        if added_conf_count > 0 or added_spk_count > 0:
            self._auto_fit_columns(ws_conf)
            self._auto_fit_columns(ws_spk)
            self._safe_save(wb)
            logger.info(f"Published {added_conf_count} conferences and {added_spk_count} speakers to Excel at {self.file_path}")

        return added_conf_count

    def read_excel_as_dataframe(self, sheet_name: str = "USA Conferences") -> pd.DataFrame:
        self._ensure_file_exists()
        return pd.read_excel(self.file_path, sheet_name=sheet_name)

    def read_all_sheets(self):
        self._ensure_file_exists()
        xls = pd.ExcelFile(self.file_path)
        result = {}
        for s in xls.sheet_names:
            result[s] = pd.read_excel(xls, sheet_name=s).fillna("").to_dict(orient="records")
        return result
